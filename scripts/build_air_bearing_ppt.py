# -*- coding: utf-8 -*-
"""生成《气浮轴承技术介绍》PPT。

用法:
    pip install python-pptx
    python scripts/build_air_bearing_ppt.py [输出路径]
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

# ---------- 主题 ----------
NAVY = RGBColor(0x0B, 0x1F, 0x3A)
NAVY_2 = RGBColor(0x12, 0x2E, 0x55)
CYAN = RGBColor(0x00, 0xB4, 0xD8)
CYAN_LIGHT = RGBColor(0xCA, 0xF0, 0xF8)
ORANGE = RGBColor(0xF7, 0x93, 0x1E)
GRAY = RGBColor(0x5F, 0x6B, 0x7A)
GRAY_LIGHT = RGBColor(0xF2, 0xF4, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x22, 0x22, 0x22)
STEEL = RGBColor(0x8D, 0x99, 0xAE)
STEEL_DARK = RGBColor(0x4A, 0x55, 0x68)

FONT = "Microsoft YaHei"

W, H = Inches(13.333), Inches(7.5)

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]
TOTAL_SLIDES = 16


# ---------- 基础工具 ----------
def set_text(tf, text, size=18, bold=False, color=BLACK, align=PP_ALIGN.LEFT, font=FONT):
    tf.text = ""
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    return p


def add_text(slide, left, top, width, height, text, size=18, bold=False, color=BLACK,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, wrap=True):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    set_text(tf, text, size, bold, color, align)
    return tb


def add_rect(slide, left, top, width, height, fill, line=None, shape=MSO_SHAPE.RECTANGLE,
             shadow=False):
    s = slide.shapes.add_shape(shape, left, top, width, height)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    if not shadow:
        s.shadow.inherit = False
    return s


def shape_text(shape, text, size=14, bold=False, color=WHITE, align=PP_ALIGN.CENTER,
               anchor=MSO_ANCHOR.MIDDLE):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    lines = text.split("\n")
    set_text(tf, lines[0], size, bold, color, align)
    for extra in lines[1:]:
        p = tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = extra
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = FONT


def add_bullets(slide, left, top, width, height, items, size=18, color=BLACK,
                bullet_color=CYAN, line_spacing=1.25, para_space=8):
    """items: list of str 或 (str, level)。"""
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    first = True
    for item in items:
        text, level = (item, 0) if isinstance(item, str) else item
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = level
        p.line_spacing = line_spacing
        p.space_after = Pt(para_space)
        marker = "■ " if level == 0 else "– "
        r0 = p.add_run()
        r0.text = ("    " * level) + marker
        r0.font.size = Pt(size - 6 if level == 0 else size - 2)
        r0.font.color.rgb = bullet_color if level == 0 else GRAY
        r0.font.name = FONT
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size if level == 0 else size - 2)
        r.font.color.rgb = color if level == 0 else GRAY
        r.font.name = FONT
    return tb


def add_line(slide, x1, y1, x2, y2, color=STEEL_DARK, width=1.5, dash=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    c.line.color.rgb = color
    c.line.width = Pt(width)
    if dash:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return c


def add_arrow(slide, left, top, width, height, color=CYAN, direction="right"):
    shape = {"right": MSO_SHAPE.RIGHT_ARROW, "down": MSO_SHAPE.DOWN_ARROW,
             "up": MSO_SHAPE.UP_ARROW, "left": MSO_SHAPE.LEFT_ARROW}[direction]
    return add_rect(slide, left, top, width, height, color, shape=shape)


# ---------- 版式 ----------
def content_slide(title, subtitle=None, index=None):
    slide = prs.slides.add_slide(BLANK)
    # 顶部标题带
    add_rect(slide, 0, 0, W, Inches(1.05), NAVY)
    add_rect(slide, 0, Inches(1.05), W, Inches(0.06), CYAN)
    add_rect(slide, Inches(0.55), Inches(0.3), Inches(0.1), Inches(0.5), CYAN)
    add_text(slide, Inches(0.8), Inches(0.2), Inches(9), Inches(0.7), title, 28, True, WHITE,
             anchor=MSO_ANCHOR.MIDDLE)
    if subtitle:
        add_text(slide, Inches(0.8), Inches(1.25), Inches(11.5), Inches(0.45), subtitle, 14,
                 False, GRAY)
    # 页脚
    add_text(slide, Inches(0.55), Inches(7.05), Inches(6), Inches(0.35), "气浮轴承技术介绍",
             11, False, STEEL)
    if index is not None:
        add_text(slide, Inches(11.3), Inches(7.05), Inches(1.5), Inches(0.35),
                 f"{index:02d} / {TOTAL_SLIDES}", 11, False, STEEL, align=PP_ALIGN.RIGHT)
    return slide


def section_slide(number, title, desc):
    slide = prs.slides.add_slide(BLANK)
    add_rect(slide, 0, 0, W, H, NAVY)
    add_rect(slide, Inches(0.9), Inches(2.6), Inches(0.12), Inches(2.1), CYAN)
    add_text(slide, Inches(1.3), Inches(2.4), Inches(4), Inches(1.2), number, 60, True, CYAN)
    add_text(slide, Inches(1.3), Inches(3.5), Inches(10), Inches(0.9), title, 40, True, WHITE)
    add_text(slide, Inches(1.3), Inches(4.4), Inches(10), Inches(0.6), desc, 18, False, STEEL)
    return slide


def card(slide, left, top, width, height, title, body, accent=CYAN, title_size=16,
         body_size=13):
    add_rect(slide, left, top, width, height, WHITE, line=RGBColor(0xDD, 0xE2, 0xEA))
    add_rect(slide, left, top, Inches(0.08), height, accent)
    add_text(slide, left + Inches(0.2), top + Inches(0.1), width - Inches(0.3), Inches(0.45),
             title, title_size, True, NAVY)
    add_text(slide, left + Inches(0.2), top + Inches(0.55), width - Inches(0.3),
             height - Inches(0.65), body, body_size, False, GRAY)


# =====================================================================
# 1. 封面
# =====================================================================
s = prs.slides.add_slide(BLANK)
add_rect(s, 0, 0, W, H, NAVY)
add_rect(s, Inches(8.6), 0, Inches(4.75), H, NAVY_2)
# 装饰：抽象气膜轴示意
cx, cy = Inches(10.95), Inches(3.75)
for i, (r, col) in enumerate([(2.4, NAVY), (1.95, CYAN), (1.75, NAVY_2), (1.2, STEEL)]):
    rr = Inches(r)
    add_rect(s, cx - rr, cy - rr, 2 * rr, 2 * rr, col, shape=MSO_SHAPE.OVAL)
add_text(s, cx - Inches(1.2), cy - Inches(0.35), Inches(2.4), Inches(0.7), "AIR\nBEARING",
         14, True, WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
add_rect(s, Inches(0.9), Inches(2.45), Inches(1.2), Inches(0.08), CYAN)
add_text(s, Inches(0.9), Inches(2.6), Inches(7.5), Inches(1.2), "气浮轴承技术介绍", 48, True,
         WHITE)
add_text(s, Inches(0.9), Inches(3.7), Inches(7.5), Inches(0.6),
         "Air Bearing Technology  |  原理 · 分类 · 性能 · 应用 · 趋势", 18, False, CYAN)
add_text(s, Inches(0.9), Inches(5.6), Inches(7.5), Inches(0.5),
         "汇报人：____________        日期：____________", 14, False, STEEL)

# =====================================================================
# 2. 目录
# =====================================================================
s = content_slide("目录", index=2)
toc = [
    ("01", "什么是气浮轴承", "定义与核心思想"),
    ("02", "工作原理", "气膜形成与承载机制"),
    ("03", "分类与结构形式", "静压 / 动压，径向 / 止推"),
    ("04", "节流方式", "小孔、多孔质、狭缝、表面节流"),
    ("05", "关键性能指标", "承载力、刚度、气膜厚度、稳定性"),
    ("06", "优势与局限", "与其他轴承的对比"),
    ("07", "典型应用", "半导体、精密机床、计量仪器等"),
    ("08", "设计要点与发展趋势", "关键技术与未来方向"),
]
for i, (num, name, desc) in enumerate(toc):
    col, row = i // 4, i % 4
    left = Inches(0.9) + col * Inches(6.2)
    top = Inches(1.6) + row * Inches(1.3)
    add_rect(s, left, top, Inches(5.7), Inches(1.05), GRAY_LIGHT)
    add_text(s, left + Inches(0.15), top + Inches(0.15), Inches(1.0), Inches(0.8), num, 28, True,
             CYAN, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, left + Inches(1.2), top + Inches(0.12), Inches(4.4), Inches(0.5), name, 18,
             True, NAVY)
    add_text(s, left + Inches(1.2), top + Inches(0.55), Inches(4.4), Inches(0.4), desc, 12,
             False, GRAY)

# =====================================================================
# 3. 什么是气浮轴承
# =====================================================================
s = content_slide("什么是气浮轴承", "Air Bearing / Gas Bearing", index=3)
add_rect(s, Inches(0.8), Inches(1.75), Inches(6.2), Inches(1.6), CYAN_LIGHT)
add_text(s, Inches(1.0), Inches(1.85), Inches(5.8), Inches(1.45),
         "气浮轴承是一种利用压缩气体（通常为空气）在相对运动表面之间形成一层极薄的"
         "高压气膜，以气膜压力支撑载荷、使两表面完全脱离接触的非接触式轴承。",
         15, False, NAVY, anchor=MSO_ANCHOR.MIDDLE)
add_bullets(s, Inches(0.8), Inches(3.6), Inches(6.2), Inches(3.3), [
    "润滑介质：空气 / 氮气等气体，而非润滑油",
    "气膜厚度：通常 5 ~ 30 μm，相当于头发丝直径的 1/10 左右",
    "运动副之间无固体接触 → 几乎零摩擦、零磨损",
    "最早可追溯至 19 世纪中叶，20 世纪 60 年代起随精密加工与航天需求快速发展",
], size=16)

# 右侧示意：轴承座 + 气膜 + 导轨
bx, by = Inches(7.6), Inches(2.0)
add_rect(s, bx, by + Inches(2.6), Inches(5.0), Inches(0.7), STEEL_DARK)          # 导轨
add_text(s, bx, by + Inches(2.7), Inches(5.0), Inches(0.5), "导轨 / 轴（运动面）", 13, True,
         WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
add_rect(s, bx + Inches(0.6), by + Inches(2.42), Inches(3.8), Inches(0.18), CYAN)  # 气膜
add_text(s, bx + Inches(4.45), by + Inches(2.3), Inches(1.0), Inches(0.4), "气膜", 12, True, CYAN)
add_rect(s, bx + Inches(0.6), by + Inches(0.6), Inches(3.8), Inches(1.8), STEEL)   # 轴承体
add_text(s, bx + Inches(0.6), by + Inches(1.0), Inches(3.8), Inches(0.6), "气浮轴承体", 15, True,
         WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
add_rect(s, bx + Inches(2.3), by, Inches(0.4), Inches(0.6), ORANGE)               # 进气口
add_text(s, bx + Inches(2.85), by - Inches(0.05), Inches(2.0), Inches(0.4), "压缩空气进气",
         12, True, ORANGE)
for k in range(4):
    ax = bx + Inches(1.0) + k * Inches(1.0)
    add_arrow(s, ax, by + Inches(1.75), Inches(0.3), Inches(0.6), ORANGE, "down")
add_text(s, bx + Inches(0.3), by + Inches(3.45), Inches(4.6), Inches(0.5),
         "气膜厚度 h ≈ 5 ~ 30 μm（放大示意）", 12, False, GRAY, align=PP_ALIGN.CENTER)

# =====================================================================
# 4. 工作原理
# =====================================================================
s = content_slide("工作原理", "从压缩气体到稳定承载的四个环节", index=4)
steps = [
    ("① 供气", "压缩机 / 气源提供洁净、干燥、稳压的压缩空气（典型 0.3 ~ 0.6 MPa）"),
    ("② 节流", "气体经节流器（小孔 / 多孔质等）降压后进入轴承间隙，节流器是产生刚度的关键"),
    ("③ 成膜", "气体在间隙内形成压力分布，向四周排出；压力积分即承载力"),
    ("④ 自平衡", "载荷↑ → 间隙↓ → 排气阻力↑ → 气膜压力↑ → 恢复平衡，形成刚度"),
]
for i, (t, d) in enumerate(steps):
    left = Inches(0.8) + i * Inches(3.1)
    add_rect(s, left, Inches(1.85), Inches(2.8), Inches(0.75), NAVY if i % 2 == 0 else CYAN,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(s, left, Inches(1.85), Inches(2.8), Inches(0.75), t, 18, True, WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, left, Inches(2.7), Inches(2.8), Inches(1.4), d, 13, False, GRAY)
    if i < 3:
        add_arrow(s, left + Inches(2.82), Inches(2.05), Inches(0.26), Inches(0.35), ORANGE)

# 压力分布示意
px, py = Inches(0.9), Inches(3.95)
add_rect(s, px, py + Inches(1.9), Inches(6.0), Inches(0.4), STEEL_DARK)
add_rect(s, px + Inches(0.5), py + Inches(1.0), Inches(5.0), Inches(0.9), STEEL)
add_rect(s, px + Inches(0.5), py + Inches(1.78), Inches(5.0), Inches(0.12), CYAN)
add_rect(s, px + Inches(2.85), py + Inches(0.55), Inches(0.3), Inches(0.45), ORANGE)
add_text(s, px + Inches(3.25), py + Inches(0.45), Inches(2.5), Inches(0.4), "节流孔 Ps", 12, True,
         ORANGE)
# 压力曲线：以倒三角近似表示间隙内的压力分布（中心高、边缘为大气压）
tri = s.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, px + Inches(0.5), py + Inches(2.3),
                         Inches(5.0), Inches(0.4))
tri.rotation = 180
tri.fill.solid()
tri.fill.fore_color.rgb = CYAN_LIGHT
tri.line.color.rgb = CYAN
add_text(s, px + Inches(0.3), py + Inches(2.72), Inches(5.4), Inches(0.35),
         "气膜压力分布 p(x)：中心高、边缘降至大气压 Pa", 11, False, NAVY, align=PP_ALIGN.CENTER)
add_text(s, px, py - Inches(0.1), Inches(6), Inches(0.4), "气膜压力分布示意", 14, True, NAVY)

add_bullets(s, Inches(7.4), Inches(4.2), Inches(5.4), Inches(2.8), [
    "承载力 W = ∫(p − Pa) dA，与供气压力、承载面积成正比",
    "刚度 K = −dW/dh，取决于节流器与间隙的匹配",
    "气体可压缩，设计不当会产生“气锤”自激振动",
    "理论基础：雷诺方程（Reynolds Equation）+ 气体状态方程",
], size=14, para_space=4)

# =====================================================================
# 5. 分类与结构形式
# =====================================================================
s = content_slide("分类与结构形式", "按气膜压力来源分为静压、动压两大类；按承载方向分为径向、止推等",
                  index=5)
card(s, Inches(0.8), Inches(1.85), Inches(5.9), Inches(2.5), "静压气浮轴承（Aerostatic）",
     "• 依靠外部气源供气，在轴承间隙内建立压力\n"
     "• 静止状态即可承载，启停无磨损\n"
     "• 刚度和承载力可通过供气压力调节\n"
     "• 需配套气源、过滤与稳压系统\n"
     "• 应用：精密机床、测量仪器、光刻机工件台", CYAN, body_size=13)
card(s, Inches(6.95), Inches(1.85), Inches(5.6), Inches(2.5), "动压气浮轴承（Aerodynamic）",
     "• 依靠表面高速相对运动将气体“楔入”间隙形成动压\n"
     "• 无需外部气源，结构简单、自给自足\n"
     "• 低速启停时存在干摩擦，需耐磨涂层\n"
     "• 典型结构：箔片轴承、螺旋槽轴承\n"
     "• 应用：涡轮增压器、空气循环机、高速鼓风机", ORANGE, body_size=13)

add_text(s, Inches(0.8), Inches(4.55), Inches(6), Inches(0.4), "按承载方向 / 结构形式", 16, True,
         NAVY)
forms = [
    ("径向轴承", "承受径向载荷\n支撑旋转主轴"),
    ("止推轴承", "承受轴向载荷\n圆盘 / 环形结构"),
    ("球面轴承", "多自由度支撑\n自动调心"),
    ("平面导轨", "直线运动导向\n气浮工作台 / 平台"),
    ("多孔质轴承", "整面均匀出气\n刚度高、抗气锤"),
]
for i, (t, d) in enumerate(forms):
    left = Inches(0.8) + i * Inches(2.4)
    add_rect(s, left, Inches(5.0), Inches(2.2), Inches(1.75), GRAY_LIGHT,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_rect(s, left + Inches(0.85), Inches(5.15), Inches(0.5), Inches(0.5), CYAN,
             shape=MSO_SHAPE.OVAL)
    add_text(s, left, Inches(5.7), Inches(2.2), Inches(0.4), t, 14, True, NAVY,
             align=PP_ALIGN.CENTER)
    add_text(s, left, Inches(6.05), Inches(2.2), Inches(0.7), d, 11, False, GRAY,
             align=PP_ALIGN.CENTER)

# =====================================================================
# 6. 节流方式
# =====================================================================
s = content_slide("节流方式", "节流器决定了静压气浮轴承的刚度、承载力与稳定性，是设计的核心",
                  index=6)
throttles = [
    ("小孔节流", "Orifice", "在轴承面加工 φ0.1 ~ 0.3 mm 的小孔，气体经小孔降压进入间隙。",
     "结构简单、成本低；刚度中等；需防止气锤振动", CYAN),
    ("多孔质节流", "Porous", "采用多孔石墨、烧结金属或陶瓷作为轴承面，气体从整个表面均匀渗出。",
     "刚度高、承载均匀、抗气锤能力强；渗透率一致性要求高", ORANGE),
    ("狭缝节流", "Slot", "通过环形或条形狭缝供气，气体沿缝隙分布进入间隙。",
     "出气均匀、承载力大；加工难度较高", NAVY),
    ("表面节流", "Surface-Restricted", "利用轴承面上极浅的均压槽或微结构进行节流。",
     "耗气量小、稳定性好；对加工精度要求极高", STEEL_DARK),
]
for i, (t, en, d, f, col) in enumerate(throttles):
    left = Inches(0.8) + i * Inches(3.05)
    top = Inches(1.85)
    add_rect(s, left, top, Inches(2.85), Inches(4.9), WHITE, line=RGBColor(0xDD, 0xE2, 0xEA))
    add_rect(s, left, top, Inches(2.85), Inches(0.9), col)
    add_text(s, left, top + Inches(0.08), Inches(2.85), Inches(0.5), t, 18, True, WHITE,
             align=PP_ALIGN.CENTER)
    add_text(s, left, top + Inches(0.5), Inches(2.85), Inches(0.35), en, 11, False, WHITE,
             align=PP_ALIGN.CENTER)
    # 迷你示意
    dx, dy = left + Inches(0.45), top + Inches(1.1)
    add_rect(s, dx, dy, Inches(1.95), Inches(0.75), STEEL)
    add_rect(s, dx, dy + Inches(0.75), Inches(1.95), Inches(0.1), CYAN)
    add_rect(s, dx, dy + Inches(0.85), Inches(1.95), Inches(0.3), STEEL_DARK)
    if i == 0:
        add_rect(s, dx + Inches(0.9), dy, Inches(0.15), Inches(0.75), ORANGE)
    elif i == 1:
        for k in range(8):
            add_rect(s, dx + Inches(0.12) + k * Inches(0.23), dy + Inches(0.55), Inches(0.08),
                     Inches(0.2), ORANGE)
    elif i == 2:
        add_rect(s, dx + Inches(0.35), dy, Inches(0.06), Inches(0.75), ORANGE)
        add_rect(s, dx + Inches(1.54), dy, Inches(0.06), Inches(0.75), ORANGE)
    else:
        add_rect(s, dx + Inches(0.9), dy, Inches(0.12), Inches(0.6), ORANGE)
        add_rect(s, dx + Inches(0.3), dy + Inches(0.6), Inches(1.35), Inches(0.1), ORANGE)
    add_text(s, left + Inches(0.15), top + Inches(2.4), Inches(2.55), Inches(1.3), d, 12,
             False, BLACK)
    add_rect(s, left + Inches(0.15), top + Inches(3.75), Inches(2.55), Inches(1.0), GRAY_LIGHT)
    add_text(s, left + Inches(0.2), top + Inches(3.8), Inches(2.45), Inches(0.9), f, 11, False,
             GRAY, anchor=MSO_ANCHOR.MIDDLE)

# =====================================================================
# 7. 关键性能指标
# =====================================================================
s = content_slide("关键性能指标", "评价气浮轴承的主要参数及影响因素", index=7)
metrics = [
    ("承载力", "Load Capacity", "气膜压力对承载面积的积分。\n影响因素：供气压力、有效承载面积、节流方式。\n典型值：0.1 ~ 0.5 MPa × 有效面积。"),
    ("刚度", "Stiffness", "单位位移所需的载荷变化量，决定抗扰动能力。\n典型值：10 ~ 500 N/μm。\n多孔质 > 狭缝 > 小孔。"),
    ("气膜厚度", "Film Thickness", "一般 5 ~ 30 μm。过厚刚度下降、耗气量增加；\n过薄易碰磨。设计常取最大刚度点对应间隙。"),
    ("耗气量", "Air Consumption", "决定气源成本与噪声。\n与间隙三次方成正比（h³），对间隙极为敏感。"),
    ("运动精度", "Motion Accuracy", "气膜的均化效应可使运动误差小于零件加工误差。\n回转精度可达 10 ~ 50 nm，直线度可达亚微米。"),
    ("稳定性", "Stability", "需避免“气锤”自激振动及高速涡动。\n措施：减小气腔容积、采用多孔质节流、增加阻尼。"),
]
for i, (t, en, d) in enumerate(metrics):
    col, row = i % 3, i // 3
    left = Inches(0.8) + col * Inches(4.05)
    top = Inches(1.85) + row * Inches(2.5)
    add_rect(s, left, top, Inches(3.85), Inches(2.3), WHITE, line=RGBColor(0xDD, 0xE2, 0xEA))
    add_rect(s, left, top, Inches(3.85), Inches(0.08), CYAN if row == 0 else ORANGE)
    add_text(s, left + Inches(0.2), top + Inches(0.18), Inches(2.2), Inches(0.45), t, 18, True,
             NAVY)
    add_text(s, left + Inches(1.9), top + Inches(0.25), Inches(1.8), Inches(0.4), en, 11, False,
             STEEL, align=PP_ALIGN.RIGHT)
    add_text(s, left + Inches(0.2), top + Inches(0.7), Inches(3.5), Inches(1.5), d, 12, False,
             GRAY)

# =====================================================================
# 8. 优势
# =====================================================================
s = content_slide("核心优势", "为什么高端装备选择气浮轴承", index=8)
adv = [
    ("极高精度", "气膜均化效应，回转误差可达纳米级，远优于滚动轴承"),
    ("零摩擦 零磨损", "非接触运行，无静摩擦、无爬行，寿命几乎无限"),
    ("高速性能", "气体黏度低，发热小，适合超高转速（可达数十万 rpm）"),
    ("洁净无污染", "无油润滑，适用于洁净室、真空、食品与医疗环境"),
    ("低振动 低噪声", "无滚动体，运行平稳，振动量级极低"),
    ("宽温度范围", "气体黏度随温度变化小，极端温度下仍可工作"),
]
for i, (t, d) in enumerate(adv):
    col, row = i % 2, i // 2
    left = Inches(0.8) + col * Inches(6.15)
    top = Inches(1.9) + row * Inches(1.6)
    add_rect(s, left, top, Inches(0.9), Inches(0.9), CYAN, shape=MSO_SHAPE.OVAL)
    add_text(s, left, top, Inches(0.9), Inches(0.9), f"{i + 1:02d}", 18, True, WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, left + Inches(1.1), top - Inches(0.05), Inches(4.8), Inches(0.5), t, 20, True,
             NAVY)
    add_text(s, left + Inches(1.1), top + Inches(0.45), Inches(4.8), Inches(0.9), d, 14, False,
             GRAY)

# =====================================================================
# 9. 局限与挑战
# =====================================================================
s = content_slide("局限与挑战", "应用气浮轴承需要面对的现实问题及应对思路", index=9)
lim = [
    ("承载力与刚度相对较低", "气体黏度仅为油的约 1/1000，同尺寸下承载能力远低于液体静压轴承。",
     "增大承载面积、提高供气压力、采用多孔质或真空预载"),
    ("依赖洁净稳定的气源", "需要压缩机、干燥器、多级过滤与稳压装置，系统复杂度与运行成本增加。",
     "配置 0.01 μm 级过滤、精密减压阀、储气罐缓冲"),
    ("加工与装配精度要求极高", "间隙仅十余微米，零件平面度、圆度需达到亚微米级。",
     "超精密磨削 / 研磨 / 刮研，恒温装配"),
    ("稳定性问题", "气体可压缩性导致气锤振动；高速动压轴承存在半速涡动。",
     "优化节流参数、减小气腔容积、增加阻尼结构"),
    ("断气风险", "一旦供气中断轴承即刻接触，可能造成划伤或抱死。",
     "设置备用气源、断气报警、耐磨涂层保护"),
]
hdr_y = Inches(1.85)
cols = [(Inches(0.8), Inches(3.2), "局限 / 挑战"), (Inches(4.1), Inches(4.6), "原因说明"),
        (Inches(8.8), Inches(3.75), "应对措施")]
for left, width, t in cols:
    add_rect(s, left, hdr_y, width, Inches(0.5), NAVY)
    add_text(s, left + Inches(0.1), hdr_y, width, Inches(0.5), t, 14, True, WHITE,
             anchor=MSO_ANCHOR.MIDDLE)
for i, (t, d, m) in enumerate(lim):
    top = hdr_y + Inches(0.5) + i * Inches(0.9)
    fill = GRAY_LIGHT if i % 2 == 0 else WHITE
    for left, width, _ in cols:
        add_rect(s, left, top, width, Inches(0.9), fill)
    add_text(s, Inches(0.9), top, Inches(3.0), Inches(0.9), t, 14, True, NAVY,
             anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(4.2), top, Inches(4.4), Inches(0.9), d, 12, False, GRAY,
             anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(8.9), top, Inches(3.55), Inches(0.9), m, 12, False, BLACK,
             anchor=MSO_ANCHOR.MIDDLE)

# =====================================================================
# 10. 与其他轴承对比
# =====================================================================
s = content_slide("与其他轴承的对比", "气浮轴承 vs 滚动轴承 vs 液体静压轴承 vs 磁悬浮轴承", index=10)
headers = ["对比项", "气浮轴承", "滚动轴承", "液体静压轴承", "磁悬浮轴承"]
rows = [
    ["摩擦 / 磨损", "几乎为零", "存在滚动摩擦与磨损", "极低（液膜）", "零"],
    ["回转精度", "★★★★★  (nm 级)", "★★☆☆☆  (μm 级)", "★★★★☆", "★★★☆☆"],
    ["承载力 / 刚度", "较低", "高", "很高", "中等"],
    ["极限转速", "很高", "中等", "中等（发热）", "很高"],
    ["洁净度", "无油，洁净", "需润滑脂 / 油", "有油污染风险", "洁净"],
    ["辅助系统", "气源 + 过滤", "无", "液压站 + 冷却", "控制器 + 传感器"],
    ["成本", "中高", "低", "高", "很高"],
    ["典型场景", "光刻机、测量机、精密主轴", "通用机械", "重型精密机床", "高速透平、飞轮储能"],
]
tbl_shape = s.shapes.add_table(len(rows) + 1, len(headers), Inches(0.8), Inches(1.85),
                               Inches(11.75), Inches(4.9))
tbl = tbl_shape.table
col_widths = [Inches(1.9), Inches(2.6), Inches(2.3), Inches(2.4), Inches(2.55)]
for j, w in enumerate(col_widths):
    tbl.columns[j].width = w
for j, htxt in enumerate(headers):
    c = tbl.cell(0, j)
    c.fill.solid()
    c.fill.fore_color.rgb = NAVY if j != 1 else CYAN
    set_text(c.text_frame, htxt, 14, True, WHITE, PP_ALIGN.CENTER)
    c.vertical_anchor = MSO_ANCHOR.MIDDLE
for i, row in enumerate(rows, start=1):
    for j, val in enumerate(row):
        c = tbl.cell(i, j)
        c.fill.solid()
        if j == 1:
            c.fill.fore_color.rgb = CYAN_LIGHT
        else:
            c.fill.fore_color.rgb = GRAY_LIGHT if i % 2 == 0 else WHITE
        set_text(c.text_frame, val, 12, j in (0, 1), NAVY if j in (0, 1) else BLACK,
                 PP_ALIGN.CENTER if j else PP_ALIGN.LEFT)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        c.margin_left = c.margin_right = Inches(0.08)

# =====================================================================
# 11. 典型应用
# =====================================================================
s = content_slide("典型应用领域", "凡是追求“极致精度、极高转速、绝对洁净”的场合，几乎都有气浮轴承的身影",
                  index=11)
apps = [
    ("半导体装备", "光刻机工件台 / 掩模台、晶圆检测、划片机主轴", CYAN),
    ("超精密机床", "金刚石车床、超精密磨床、非球面加工主轴", NAVY),
    ("精密计量仪器", "三坐标测量机（CMM）、圆度仪、轮廓仪、激光干涉台", ORANGE),
    ("高速旋转机械", "涡轮增压器、空气循环机、高速电主轴、牙科手机（50 万 rpm）", STEEL_DARK),
    ("航空航天", "惯性导航陀螺、卫星姿控飞轮、低温透平膨胀机", CYAN),
    ("电子与信息", "硬盘磁头（动压气膜）、PCB 钻孔主轴、晶圆搬运机器人", NAVY),
    ("医疗与科研", "CT 旋转机架、显微镜载物台、同步辐射光束线精密平台", ORANGE),
    ("能源与低温", "氦气透平、LNG 膨胀机、燃料电池空压机", STEEL_DARK),
]
for i, (t, d, col) in enumerate(apps):
    c_, r_ = i % 4, i // 4
    left = Inches(0.8) + c_ * Inches(3.05)
    top = Inches(1.85) + r_ * Inches(2.5)
    add_rect(s, left, top, Inches(2.85), Inches(2.3), WHITE, line=RGBColor(0xDD, 0xE2, 0xEA))
    add_rect(s, left, top, Inches(2.85), Inches(0.75), col)
    add_text(s, left, top, Inches(2.85), Inches(0.75), t, 17, True, WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, left + Inches(0.15), top + Inches(0.9), Inches(2.55), Inches(1.3), d, 13, False,
             GRAY)

# =====================================================================
# 12. 应用案例聚焦：光刻机工件台
# =====================================================================
s = content_slide("应用聚焦：光刻机气浮工件台", "气浮轴承在超精密运动系统中的典型集成方案", index=12)
# 示意图
gx, gy = Inches(0.9), Inches(2.0)
add_rect(s, gx, gy + Inches(3.3), Inches(6.0), Inches(0.9), STEEL_DARK)
add_text(s, gx, gy + Inches(3.3), Inches(6.0), Inches(0.9), "花岗岩基座（平面度 < 1 μm）", 13, True,
         WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
add_rect(s, gx + Inches(1.0), gy + Inches(3.12), Inches(4.0), Inches(0.18), CYAN)
add_rect(s, gx + Inches(1.0), gy + Inches(1.9), Inches(4.0), Inches(1.22), STEEL)
add_text(s, gx + Inches(1.0), gy + Inches(1.9), Inches(4.0), Inches(1.22),
         "气浮滑块 / 工件台\n（底部多个气浮垫 + 真空预载）", 13, True, WHITE, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE)
add_rect(s, gx + Inches(1.6), gy + Inches(1.3), Inches(2.8), Inches(0.6), ORANGE)
add_text(s, gx + Inches(1.6), gy + Inches(1.3), Inches(2.8), Inches(0.6), "晶圆 / 掩模", 13, True,
         WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
add_arrow(s, gx + Inches(0.1), gy + Inches(2.3), Inches(0.8), Inches(0.4), CYAN, "left")
add_arrow(s, gx + Inches(5.1), gy + Inches(2.3), Inches(0.8), Inches(0.4), CYAN, "right")
add_text(s, gx, gy + Inches(0.6), Inches(6.0), Inches(0.5), "直线电机驱动 + 激光干涉仪 / 光栅反馈",
         13, True, NAVY, align=PP_ALIGN.CENTER)
add_text(s, gx + Inches(5.1), gy + Inches(3.05), Inches(1.2), Inches(0.35), "气膜", 11, True, CYAN)

add_bullets(s, Inches(7.3), Inches(1.85), Inches(5.3), Inches(5.0), [
    "需求：纳米级定位精度 + 米/秒级高速扫描 + 洁净无尘",
    "方案：静压气浮垫承载，真空预载提高刚度、消除间隙漂移",
    "导向：花岗岩导轨 + 多组气浮垫形成闭环约束",
    "驱动：直线电机非接触驱动，与气浮轴承配合实现“全非接触”运动",
    "反馈：双频激光干涉仪，分辨率 < 1 nm",
    "效果：重复定位精度 ±几 nm，无磨损、无爬行、使用寿命长",
], size=14, para_space=6)

# =====================================================================
# 13. 设计要点与关键技术
# =====================================================================
s = content_slide("设计要点与关键技术", "一套可靠的气浮轴承系统 = 结构设计 + 加工工艺 + 气源保障 + 仿真验证",
                  index=13)
keys = [
    ("结构参数设计", ["间隙、节流孔径 / 数量、气腔容积的匹配优化", "最大刚度点设计，兼顾耗气量",
                "预载方式：重力 / 真空 / 磁力 / 对置轴承"]),
    ("材料选择", ["轴承体：铝合金、不锈钢、陶瓷、多孔石墨", "导轨：花岗岩、陶瓷、硬质合金",
              "考虑热膨胀匹配与耐磨性"]),
    ("超精密加工", ["平面度 / 圆度 ≤ 0.5 μm，表面粗糙度 Ra ≤ 0.1 μm", "小孔微细加工（激光 / 电火花）",
               "多孔质材料渗透率的均匀性控制"]),
    ("气源系统", ["多级过滤（≤ 0.01 μm）+ 冷冻 / 吸附干燥", "精密减压稳压，压力波动 < 1%",
              "储气罐缓冲、断气保护"]),
    ("仿真与试验", ["CFD / 雷诺方程求解压力分布", "刚度、承载、稳定性（气锤）分析",
               "电容测微 / 激光干涉测试气膜厚度与刚度"]),
    ("集成与控制", ["与直线电机、光栅 / 干涉仪的集成", "热管理与环境控制（恒温 ±0.1 ℃）",
               "主动控制气浮轴承（可变节流）"]),
]
for i, (t, items) in enumerate(keys):
    col, row = i % 3, i // 3
    left = Inches(0.8) + col * Inches(4.05)
    top = Inches(1.85) + row * Inches(2.5)
    add_rect(s, left, top, Inches(3.85), Inches(2.3), GRAY_LIGHT,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_rect(s, left + Inches(0.2), top + Inches(0.2), Inches(0.12), Inches(0.4), CYAN)
    add_text(s, left + Inches(0.45), top + Inches(0.12), Inches(3.2), Inches(0.5), t, 17, True,
             NAVY)
    add_bullets(s, left + Inches(0.3), top + Inches(0.7), Inches(3.4), Inches(1.6), items,
                size=12, para_space=2, line_spacing=1.1)

# =====================================================================
# 14. 发展趋势
# =====================================================================
s = content_slide("发展趋势", "气浮轴承技术的未来方向", index=14)
trends = [
    ("更高刚度", "多孔质 + 真空预载、复合节流、主动控制气浮轴承，突破气浮刚度瓶颈"),
    ("智能化", "内置压力 / 位移传感器，实时监测气膜状态，实现故障预警与自适应调节"),
    ("新材料", "多孔陶瓷、碳纤维复合材料、新型多孔石墨，提升刚度、减轻重量、降低热变形"),
    ("集成化 / 模块化", "气浮轴承与直线电机、光栅尺一体化设计，形成标准化气浮运动模组"),
    ("动压箔片轴承", "面向无油透平机械（燃料电池空压机、微型燃气轮机）的高温高速箔片轴承"),
    ("数字化设计", "多物理场耦合仿真、数字孪生与 AI 优化，缩短研发周期"),
]
add_line(s, Inches(1.4), Inches(2.3), Inches(1.4), Inches(6.7), CYAN, 3)
for i, (t, d) in enumerate(trends):
    top = Inches(2.0) + i * Inches(0.78)
    add_rect(s, Inches(1.22), top + Inches(0.12), Inches(0.36), Inches(0.36), ORANGE,
             shape=MSO_SHAPE.OVAL)
    add_text(s, Inches(1.8), top, Inches(3.0), Inches(0.6), t, 17, True, NAVY,
             anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(4.7), top, Inches(8.0), Inches(0.6), d, 13, False, GRAY,
             anchor=MSO_ANCHOR.MIDDLE)

# =====================================================================
# 15. 总结
# =====================================================================
s = content_slide("总结", index=15)
summary = [
    ("原理", "压缩气体在微米级间隙内形成气膜，以压力支撑载荷，实现完全非接触运动。"),
    ("分类", "静压（外部供气）与动压（自生压力）两大类；径向、止推、平面导轨等多种结构。"),
    ("特点", "极高精度、零摩擦零磨损、高速洁净；但承载刚度较低、依赖气源、加工要求高。"),
    ("应用", "半导体光刻、超精密机床、精密计量、高速旋转机械、航空航天等高端领域的关键部件。"),
    ("趋势", "高刚度化、智能化、新材料、模块化集成与数字化设计。"),
]
for i, (t, d) in enumerate(summary):
    top = Inches(1.7) + i * Inches(1.0)
    add_rect(s, Inches(0.9), top, Inches(1.6), Inches(0.8), NAVY if i % 2 == 0 else CYAN,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(s, Inches(0.9), top, Inches(1.6), Inches(0.8), t, 18, True, WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, Inches(2.6), top, Inches(9.9), Inches(0.8), GRAY_LIGHT)
    add_text(s, Inches(2.8), top, Inches(9.6), Inches(0.8), d, 15, False, BLACK,
             anchor=MSO_ANCHOR.MIDDLE)

# =====================================================================
# 16. 结束页
# =====================================================================
s = prs.slides.add_slide(BLANK)
add_rect(s, 0, 0, W, H, NAVY)
add_rect(s, Inches(5.5), Inches(2.3), Inches(2.3), Inches(0.08), CYAN)
add_text(s, 0, Inches(2.6), W, Inches(1.2), "谢谢观看", 54, True, WHITE, align=PP_ALIGN.CENTER)
add_text(s, 0, Inches(3.8), W, Inches(0.6), "Thank You  |  欢迎交流与指正", 20, False, CYAN,
         align=PP_ALIGN.CENTER)
add_text(s, 0, Inches(5.6), W, Inches(0.5), "气浮轴承技术介绍", 14, False, STEEL,
         align=PP_ALIGN.CENTER)

# ---------- 输出 ----------
out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parent.parent / "气浮轴承技术介绍.pptx"
out.parent.mkdir(parents=True, exist_ok=True)
prs.save(str(out))
print(f"saved: {out}  ({len(prs.slides)} slides)")
